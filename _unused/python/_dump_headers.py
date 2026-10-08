import openpyxl
from pathlib import Path

SRC = Path(r"C:\2026_Projects_Ongoing\ETL_Syndicationx_Code\06July2026_ETL\python\mapFolder\Amberwood Apartment Homes\Cash_Flow_Portal_Investments_07_17_2026.xlsx")
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
print("sheets", wb.sheetnames)
ws = wb["Investors (profile)"]
rows = ws.iter_rows(values_only=True)
headers = [str(h or "").strip() for h in next(rows)]
for i, h in enumerate(headers):
    print(f"{i:02d}: {h}")
print("--- first data row ---")
row1 = next(rows)
for i, h in enumerate(headers):
    v = row1[i] if i < len(row1) else None
    if v is not None and str(v).strip() != "":
        print(f"{i:02d} {h} = {v!r}")
wb.close()
