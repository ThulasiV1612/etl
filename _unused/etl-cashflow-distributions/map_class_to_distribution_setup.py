"""
Class_Mapping_FILLED.xlsx → ETL_Distribution_Setup_FILLED.xlsx + .payload.json

Same pattern as etl-cashflow-classes / etl-cashflow-profile-investors:
  1. Drop Class_Mapping_FILLED.xlsx into ./input/ (or pass a path)
  2. Run this script
  3. Outputs land in ./output/
  4. Load with load_distribution_setup.py (resolve class names → UUIDs, PUT API)

Mapping rules (from Class Setup → waterfall config):
  - Each LP with pref_enabled → operating LP_PREF → that LP class
  - Each preferred_equity → operating PREF_CURRENT + capital PREF_ACCRUED
  - Each LP → capital ROC → that LP class
  - Each GP → capital CATCHUP → that GP class (final % from Final Tier)

Usage:
  python map_class_to_distribution_setup.py
  python map_class_to_distribution_setup.py path/to/Class_Mapping_FILLED.xlsx
  python map_class_to_distribution_setup.py --all-migration-deals
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

import openpyxl

from dist_workbook import (
    DistSetupMapped,
    WaterfallRow,
    slug,
    write_blank_template,
    write_filled_workbook,
    write_payload,
)

HERE = Path(__file__).resolve().parent
INPUT_DIR = HERE / "input"
OUTPUT_DIR = HERE / "output"
TEMPLATES = HERE / "templates"
TEMPLATE_BLANK = TEMPLATES / "ETL_Distribution_Setup.xlsx"
REPO = HERE.parent

MIGRATION_SOURCES = {
    "Horizon Apartments": [
        REPO / "Horizon" / "Class_Mapping_FILLED.xlsx",
        REPO / "Investor-Portal-main" / "Class_Mapping_FILLED.xlsx",
    ],
    "Wildflower Assisted Living": [
        REPO / "Wildflower" / "Class_Mapping_FILLED.xlsx",
    ],
    "Woodland Ridge": [
        REPO / "Woodland_Ridge" / "Class_Mapping_FILLED.xlsx",
    ],
}

# Also copy FILLED artifacts next to deal folders (same as prior scripts)
DEAL_OUTPUT_COPY = {
    "Horizon Apartments": REPO / "Investor-Portal-main",
    "Wildflower Assisted Living": REPO / "Wildflower",
    "Woodland Ridge": REPO / "Woodland_Ridge",
}


def read_sheet(path: Path, name: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if name not in wb.sheetnames:
        wb.close()
        raise KeyError(f"Sheet {name!r} not in {path.name}: {wb.sheetnames}")
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    out: list[dict] = []
    for row in rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        first = str(row[0] or "").strip().upper()
        if first in {"LOOKUP", "HINT"}:
            continue
        out.append(
            {
                headers[i]: ("" if row[i] is None else str(row[i]).strip())
                for i in range(len(headers))
                if headers[i]
            }
        )
    wb.close()
    return out


def yesish(v: str) -> bool:
    return (v or "").strip().lower() in {"yes", "true", "1", "y"}


def sub_type(c: dict) -> str:
    return (c.get("subscription_type") or c.get("class_type") or "").strip().lower()


def build_from_class_mapping(path: Path) -> DistSetupMapped:
    classes = read_sheet(path, "Classes")
    if not classes:
        raise ValueError(f"No class rows in {path}")
    try:
        hurdles = read_sheet(path, "Hurdles")
    except KeyError:
        hurdles = []
    try:
        final = read_sheet(path, "Final Tier")
    except KeyError:
        final = []

    deal = (classes[0].get("deal_name") or "").strip() or path.stem
    deal_id = (classes[0].get("deal_id") or "").strip()
    target = (classes[0].get("deal_target_raise") or "0").strip() or "0"

    lps = [c for c in classes if sub_type(c) == "lp"]
    gps = [c for c in classes if sub_type(c) == "gp"]
    prefs = [c for c in classes if sub_type(c) == "preferred_equity"]

    operating: list[WaterfallRow] = []
    capital: list[WaterfallRow] = []

    for c in lps:
        if yesish(c.get("pref_enabled") or ""):
            rate = (c.get("pref_rate") or "").strip()
            name = (
                f"LP preferred return (+ arrears) - {rate}%"
                if rate
                else "LP preferred return (+ arrears)"
            )
            operating.append(
                WaterfallRow(
                    deal_name=deal,
                    deal_id=deal_id,
                    row_order=len(operating) + 1,
                    row_id=f"op_lp_pref_{slug(c.get('name') or 'lp')}",
                    kind="LP_PREF",
                    name=name,
                    pay_to_class_name=c.get("name") or "",
                    source=f"Classes.pref_rate={rate}",
                )
            )

    for c in prefs:
        cname = c.get("name") or ""
        operating.append(
            WaterfallRow(
                deal_name=deal,
                deal_id=deal_id,
                row_order=len(operating) + 1,
                row_id=f"op_pref_current_{slug(cname)}",
                kind="PREF_CURRENT",
                name="Preferred equity - current pay",
                pay_to_class_name=cname,
                source="preferred_equity",
            )
        )
        capital.append(
            WaterfallRow(
                deal_name=deal,
                deal_id=deal_id,
                row_order=len(capital) + 1,
                row_id=f"cap_pref_accrued_{slug(cname)}",
                kind="PREF_ACCRUED",
                name="Preferred equity - accrued balance",
                pay_to_class_name=cname,
                source="preferred_equity",
            )
        )

    for c in lps:
        capital.append(
            WaterfallRow(
                deal_name=deal,
                deal_id=deal_id,
                row_order=len(capital) + 1,
                row_id=f"cap_roc_{slug(c.get('name') or 'lp')}",
                kind="ROC",
                name="Return of LP capital",
                pay_to_class_name=c.get("name") or "",
                source="LP class",
            )
        )

    ft = final[0] if final else {}
    final_lp = ft.get("final_lp_pct") or "50"
    final_gp = ft.get("final_gp_pct") or "50"
    for c in gps:
        capital.append(
            WaterfallRow(
                deal_name=deal,
                deal_id=deal_id,
                row_order=len(capital) + 1,
                row_id=f"cap_catchup_{slug(c.get('name') or 'gp')}",
                kind="CATCHUP",
                name=f"GP catch-up (final {final_lp}/{final_gp})",
                pay_to_class_name=c.get("name") or "",
                catchup_pct="20",
                source="GP + Final Tier",
            )
        )

    return DistSetupMapped(
        deal_name=deal,
        deal_id=deal_id,
        target_raise=target,
        mapped_from=path.name,
        operating=operating,
        capital=capital,
        class_setup_context={
            "classes": [
                {
                    "name": c.get("name"),
                    "type": sub_type(c),
                    "pref": c.get("pref_rate"),
                    "equity": c.get("equity_pct"),
                }
                for c in classes
            ],
            "hurdles": hurdles,
            "final": final,
        },
    )


def resolve_input(explicit: str | None) -> Path:
    if explicit:
        p = Path(explicit).expanduser().resolve()
        if not p.exists():
            raise FileNotFoundError(p)
        return p
    files = [
        f
        for f in INPUT_DIR.glob("*.xlsx")
        if not f.name.startswith("~") and "FILLED" in f.name.upper()
    ]
    if not files:
        files = [f for f in INPUT_DIR.glob("Class_Mapping*.xlsx") if not f.name.startswith("~")]
    if not files:
        raise FileNotFoundError(
            f"No Class_Mapping_FILLED.xlsx in {INPUT_DIR}. "
            "Pass a path or drop the file into input/"
        )
    if len(files) > 1:
        newest = max(files, key=lambda f: f.stat().st_mtime)
        print("multiple input xlsx; using newest:", newest.name)
        return newest
    return files[0]


def ensure_dirs() -> None:
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    TEMPLATES.mkdir(parents=True, exist_ok=True)
    if not TEMPLATE_BLANK.exists():
        write_blank_template(TEMPLATE_BLANK)


def process_one(src: Path, out_dir: Path, *, also_copy_to: Path | None = None) -> DistSetupMapped:
    mapped = build_from_class_mapping(src)
    safe = slug(mapped.deal_name) or "deal"
    xlsx = out_dir / f"ETL_Distribution_Setup_FILLED_{safe}.xlsx"
    payload = out_dir / f"ETL_Distribution_Setup_FILLED_{safe}.payload.json"
    # Canonical names inside out_dir when single-deal
    xlsx_std = out_dir / "ETL_Distribution_Setup_FILLED.xlsx"
    payload_std = out_dir / "ETL_Distribution_Setup_FILLED.payload.json"

    write_filled_workbook(mapped, xlsx)
    write_payload(mapped, payload)
    write_filled_workbook(mapped, xlsx_std)
    write_payload(mapped, payload_std)

    if also_copy_to:
        also_copy_to.mkdir(parents=True, exist_ok=True)
        if mapped.deal_name == "Horizon Apartments":
            shutil.copy2(xlsx, also_copy_to / "ETL_Distribution_Setup_FILLED_Horizon.xlsx")
            shutil.copy2(payload, also_copy_to / "ETL_Distribution_Setup_FILLED_Horizon.payload.json")
        else:
            shutil.copy2(xlsx_std, also_copy_to / "ETL_Distribution_Setup_FILLED.xlsx")
            shutil.copy2(payload_std, also_copy_to / "ETL_Distribution_Setup_FILLED.payload.json")
        # Root convenience copies for Wildflower / Woodland naming
        if mapped.deal_name.startswith("Wildflower"):
            shutil.copy2(payload, REPO / "ETL_Distribution_Setup_FILLED_Wildflower.payload.json")
        elif mapped.deal_name.startswith("Woodland"):
            shutil.copy2(payload, REPO / "ETL_Distribution_Setup_FILLED_Woodland_Ridge.payload.json")

    print(f"source={src}")
    print(f"deal={mapped.deal_name} operating={len(mapped.operating)} capital={len(mapped.capital)}")
    print(f"wrote={xlsx.name}")
    print(f"wrote={payload.name}")
    return mapped


def find_migration_source(deal_name: str) -> Path:
    for p in MIGRATION_SOURCES.get(deal_name, []):
        if p.exists():
            return p
    raise FileNotFoundError(f"No Class_Mapping_FILLED for {deal_name}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Map Class_Mapping_FILLED → ETL Distribution Setup FILLED + payload"
    )
    parser.add_argument(
        "class_mapping",
        nargs="?",
        help="Class_Mapping_FILLED.xlsx. Default: newest FILLED xlsx in ./input/",
    )
    parser.add_argument("--out-dir", default=str(OUTPUT_DIR))
    parser.add_argument(
        "--all-migration-deals",
        action="store_true",
        help="Map Horizon, Wildflower, Woodland from repo Class_Mapping_FILLED files",
    )
    parser.add_argument(
        "--no-copy",
        action="store_true",
        help="Do not copy artifacts into deal folders",
    )
    args = parser.parse_args(argv)

    ensure_dirs()
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.all_migration_deals:
        for deal_name in MIGRATION_SOURCES:
            src = find_migration_source(deal_name)
            deal_out = out_dir / slug(deal_name)
            deal_out.mkdir(parents=True, exist_ok=True)
            copy_to = None if args.no_copy else DEAL_OUTPUT_COPY.get(deal_name)
            print(f"\n=== {deal_name} ===")
            process_one(src, deal_out, also_copy_to=copy_to)
        print("\nDone — all migration deals.")
        return 0

    src = resolve_input(args.class_mapping)
    copy_to = None
    if not args.no_copy:
        # Best-effort: match deal after parse
        mapped_preview = build_from_class_mapping(src)
        copy_to = DEAL_OUTPUT_COPY.get(mapped_preview.deal_name)
    process_one(src, out_dir, also_copy_to=copy_to)
    print("Done.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        raise SystemExit(1)
