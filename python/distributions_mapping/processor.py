"""
Class_Mapping_FILLED.xlsx → ETL_Distribution_Setup_FILLED.xlsx + .payload.json

Writes under mapFolder/<deal_name>/ (same pattern as profile/classes mappers).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

from dist_workbook import (
    DistSetupMapped,
    WaterfallRow,
    slug,
    write_blank_template,
    write_filled_workbook,
    write_payload,
)

HERE = Path(__file__).resolve().parent
TEMPLATES = HERE / "templates"
TEMPLATE_BLANK = TEMPLATES / "ETL_Distribution_Setup.xlsx"

OUT_XLSX = "ETL_Distribution_Setup_FILLED.xlsx"
OUT_PAYLOAD = "ETL_Distribution_Setup_FILLED.payload.json"
OUT_WARNINGS = "warnings.txt"


def safe_folder_name(name: str) -> str:
    raw = str(name or "").strip() or "Unknown_Deal"
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", raw)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    cleaned = cleaned[:120] if cleaned else "Unknown_Deal"
    return cleaned or "Unknown_Deal"


def read_sheet(path: Path, name: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if name not in wb.sheetnames:
        wb.close()
        raise KeyError(f"Sheet {name!r} not in {path.name}: {wb.sheetnames}")
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    out: list[dict] = []
    for row in rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        first = str(row[0] or "").strip().upper()
        if first in {"LOOKUP", "HINT"}:
            continue
        out.append(
            {
                headers[i]: ("" if row[i] is None else str(row[i]).strip())
                for i in range(len(headers))
                if headers[i]
            }
        )
    wb.close()
    return out


def yesish(v: str) -> bool:
    return (v or "").strip().lower() in {"yes", "true", "1", "y"}


def sub_type(c: dict) -> str:
    return (c.get("subscription_type") or c.get("class_type") or "").strip().lower()


def build_from_class_mapping(path: Path) -> DistSetupMapped:
    try:
        classes = read_sheet(path, "Classes")
    except KeyError as exc:
        raise ValueError("Classes sheet not found. Upload a Class_Mapping_FILLED.xlsx file.") from exc

    if not classes:
        raise ValueError("No class rows found in the Class Mapping workbook.")

    try:
        hurdles = read_sheet(path, "Hurdles")
    except KeyError:
        hurdles = []
    try:
        final = read_sheet(path, "Final Tier")
    except KeyError:
        final = []

    deal = (classes[0].get("deal_name") or "").strip() or path.stem
    deal_id = (classes[0].get("deal_id") or "").strip()
    target = (classes[0].get("deal_target_raise") or "0").strip() or "0"

    lps = [c for c in classes if sub_type(c) == "lp"]
    gps = [c for c in classes if sub_type(c) == "gp"]
    prefs = [c for c in classes if sub_type(c) == "preferred_equity"]

    operating: list[WaterfallRow] = []
    capital: list[WaterfallRow] = []

    for c in lps:
        if yesish(c.get("pref_enabled") or ""):
            rate = (c.get("pref_rate") or "").strip()
            name = (
                f"LP preferred return (+ arrears) - {rate}%"
                if rate
                else "LP preferred return (+ arrears)"
            )
            operating.append(
                WaterfallRow(
                    deal_name=deal,
                    deal_id=deal_id,
                    row_order=len(operating) + 1,
                    row_id=f"op_lp_pref_{slug(c.get('name') or 'lp')}",
                    kind="LP_PREF",
                    name=name,
                    pay_to_class_name=c.get("name") or "",
                    source=f"Classes.pref_rate={rate}",
                )
            )

    for c in prefs:
        cname = c.get("name") or ""
        operating.append(
            WaterfallRow(
                deal_name=deal,
                deal_id=deal_id,
                row_order=len(operating) + 1,
                row_id=f"op_pref_current_{slug(cname)}",
                kind="PREF_CURRENT",
                name="Preferred equity - current pay",
                pay_to_class_name=cname,
                source="preferred_equity",
            )
        )
        capital.append(
            WaterfallRow(
                deal_name=deal,
                deal_id=deal_id,
                row_order=len(capital) + 1,
                row_id=f"cap_pref_accrued_{slug(cname)}",
                kind="PREF_ACCRUED",
                name="Preferred equity - accrued balance",
                pay_to_class_name=cname,
                source="preferred_equity",
            )
        )

    for c in lps:
        capital.append(
            WaterfallRow(
                deal_name=deal,
                deal_id=deal_id,
                row_order=len(capital) + 1,
                row_id=f"cap_roc_{slug(c.get('name') or 'lp')}",
                kind="ROC",
                name="Return of LP capital",
                pay_to_class_name=c.get("name") or "",
                source="LP class",
            )
        )

    ft = final[0] if final else {}
    final_lp = ft.get("final_lp_pct") or "50"
    final_gp = ft.get("final_gp_pct") or "50"
    for c in gps:
        capital.append(
            WaterfallRow(
                deal_name=deal,
                deal_id=deal_id,
                row_order=len(capital) + 1,
                row_id=f"cap_catchup_{slug(c.get('name') or 'gp')}",
                kind="CATCHUP",
                name=f"GP catch-up (final {final_lp}/{final_gp})",
                pay_to_class_name=c.get("name") or "",
                catchup_pct="20",
                source="GP + Final Tier",
            )
        )

    return DistSetupMapped(
        deal_name=deal,
        deal_id=deal_id,
        target_raise=target,
        mapped_from=path.name,
        operating=operating,
        capital=capital,
        class_setup_context={
            "classes": [
                {
                    "name": c.get("name"),
                    "type": sub_type(c),
                    "pref": c.get("pref_rate"),
                    "equity": c.get("equity_pct"),
                }
                for c in classes
            ],
            "hurdles": hurdles,
            "final": final,
        },
    )


def apply_deal_id(mapped: DistSetupMapped, deal_id: str | None) -> DistSetupMapped:
    override = str(deal_id or "").strip()
    if not override:
        return mapped
    mapped.deal_id = override
    for row in mapped.operating:
        row.deal_id = override
    for row in mapped.capital:
        row.deal_id = override
    return mapped


def process_workbook(
    input_path: Path,
    out_dir: Path,
    deal_id: str | None = None,
) -> dict:
    """Write distribution setup files under out_dir/<deal_name>/."""
    if not input_path.exists():
        raise FileNotFoundError(input_path)
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("Please upload an Excel (.xlsx) Class Mapping file.")

    TEMPLATES.mkdir(parents=True, exist_ok=True)
    if not TEMPLATE_BLANK.exists():
        write_blank_template(TEMPLATE_BLANK)

    map_root = Path(out_dir).expanduser().resolve()
    map_root.mkdir(parents=True, exist_ok=True)

    mapped = apply_deal_id(build_from_class_mapping(input_path), deal_id)
    warnings: list[str] = []

    if not mapped.operating and not mapped.capital:
        raise RuntimeError(
            "No distribution waterfall rows could be generated from the Class Mapping file."
        )

    if not mapped.operating:
        warnings.append("No operating waterfall rows were generated.")
    if not mapped.capital:
        warnings.append("No capital waterfall rows were generated.")
    if not mapped.deal_id:
        warnings.append("deal_id is empty — resolve deal UUID before loading to the database.")

    folder_name = safe_folder_name(mapped.deal_name)
    deal_dir = map_root / folder_name
    deal_dir.mkdir(parents=True, exist_ok=True)

    xlsx_path = deal_dir / OUT_XLSX
    payload_path = deal_dir / OUT_PAYLOAD
    warnings_path = deal_dir / OUT_WARNINGS

    write_filled_workbook(mapped, xlsx_path)
    write_payload(mapped, payload_path)

    if warnings:
        warnings_path.write_text("\n".join(warnings) + "\n", encoding="utf-8")
    elif warnings_path.exists():
        warnings_path.unlink()

    return {
        "success": True,
        "sourceFile": input_path.name,
        "summary": {
            "dealName": mapped.deal_name,
            "dealId": mapped.deal_id or "",
            "operatingRows": len(mapped.operating),
            "capitalRows": len(mapped.capital),
            "targetRaise": mapped.target_raise or "",
            "warnings": len(warnings),
        },
        "warnings": warnings,
        "dealFolders": [
            {
                "dealName": mapped.deal_name,
                "folderName": folder_name,
                "files": {
                    "workbook": str(xlsx_path),
                    "payload": str(payload_path),
                    "warnings": str(warnings_path) if warnings else None,
                },
            }
        ],
        "files": {
            "workbook": str(xlsx_path),
            "payload": str(payload_path),
            "warnings": str(warnings_path) if warnings else None,
        },
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Map Class_Mapping_FILLED.xlsx to Distribution Setup FILLED files"
    )
    parser.add_argument("class_mapping", help="Class_Mapping_FILLED.xlsx path")
    parser.add_argument("--out-dir", required=True, help="mapFolder root")
    parser.add_argument("--deal-id", default="", help="Override deal UUID in output")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    try:
        result = process_workbook(
            Path(args.class_mapping).expanduser().resolve(),
            Path(args.out_dir).expanduser().resolve(),
            deal_id=args.deal_id or None,
        )
    except Exception as e:
        message = str(e)
        if "Classes sheet not found" in message:
            message = "Classes sheet not found. Upload a Class_Mapping_FILLED.xlsx file."
        elif "No class rows" in message:
            message = "No class rows found in the Class Mapping workbook."
        print(json.dumps({"success": False, "error": message}), file=sys.stderr)
        print("ERROR:", message, file=sys.stderr)
        return 1

    if args.json:
        print(json.dumps(result))
    else:
        print("deal", result["summary"]["dealName"])
        print("operating", result["summary"]["operatingRows"])
        print("capital", result["summary"]["capitalRows"])
        print("wrote", result["files"]["workbook"])
        print("wrote", result["files"]["payload"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
