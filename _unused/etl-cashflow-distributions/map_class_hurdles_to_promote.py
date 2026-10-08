"""
Class_Mapping_FILLED Hurdles / Final Tier → Class Setup promote schedule payload.

UI Class Setup "Hurdles & promote schedule" reads meta.promote, not only
per-class waterfallTiers. This script builds promote so rates (e.g. 7%) show.

Usage:
  python map_class_hurdles_to_promote.py
  python map_class_hurdles_to_promote.py path/to/Class_Mapping_FILLED.xlsx
  python map_class_hurdles_to_promote.py --all-migration-deals
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
INPUT_DIR = HERE / "input"
OUTPUT_DIR = HERE / "output"
REPO = HERE.parent

MIGRATION_SOURCES = {
    "Horizon Apartments": [
        REPO / "Horizon" / "Class_Mapping_FILLED.xlsx",
        REPO / "distributions" / "Class_Mapping_FILLED.xlsx",
        REPO / "Investor-Portal-main" / "Class_Mapping_FILLED.xlsx",
    ],
    "Wildflower Assisted Living": [REPO / "Wildflower" / "Class_Mapping_FILLED.xlsx"],
    "Woodland Ridge": [REPO / "Woodland_Ridge" / "Class_Mapping_FILLED.xlsx"],
}


def cell(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def read_sheet(path: Path, name: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if name not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    headers = [cell(h) for h in next(rows)]
    out: list[dict] = []
    for row in rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        if cell(row[0]).upper() in {"LOOKUP", "HINT"}:
            continue
        out.append({headers[i]: cell(row[i]) if i < len(row) else "" for i in range(len(headers)) if headers[i]})
    wb.close()
    return out


def basis_from_pref_type(raw: str) -> str:
    t = (raw or "").lower()
    if "irr" in t:
        return "IRR"
    if "cash" in t or "coc" in t:
        return "Cash-on-cash"
    return "Cash-on-cash"


def build_promote(path: Path) -> dict:
    classes = read_sheet(path, "Classes")
    hurdles = read_sheet(path, "Hurdles")
    final = read_sheet(path, "Final Tier")
    deal = (classes[0].get("deal_name") if classes else "") or path.stem

    promote_hurdles = []
    seen = set()
    for h in hurdles:
        if (h.get("is_final") or "").lower() in {"yes", "true", "1"}:
            continue
        rate = (h.get("hurdle_rate_pct") or "").replace("%", "").strip()
        if not rate or rate in seen:
            continue
        seen.add(rate)
        promote_hurdles.append(
            {
                "id": f"h{len(promote_hurdles) + 1}",
                "rate": rate,
                "basis": basis_from_pref_type(h.get("preferred_return_type") or ""),
                "measuredOn": "LP classes",
                "source_class_name": h.get("class_name") or "",
                "lp_pct": (h.get("lp_pct") or "").replace("%", ""),
                "gp_pct": (h.get("gp_pct") or "").replace("%", ""),
            }
        )

    ft = final[0] if final else {}
    final_lp = (ft.get("final_lp_pct") or "50").replace("%", "")
    final_gp = (ft.get("final_gp_pct") or "50").replace("%", "")

    # shares keyed by class NAME until load resolves to class UUIDs
    shares_by_name: dict[str, list[str]] = {}
    stage_n = max(1, len(promote_hurdles)) + 1
    for c in classes:
        name = c.get("name") or ""
        if not name:
            continue
        st = (c.get("subscription_type") or "").lower()
        if st == "lp":
            # match first hurdle for this class if present
            match = next((h for h in hurdles if h.get("class_name") == name), None)
            lp = (match.get("lp_pct") if match else "") or "100"
            lp = lp.replace("%", "")
            shares_by_name[name] = [
                lp if i < stage_n - 1 else final_lp for i in range(stage_n)
            ]
        elif st == "gp":
            match = next((h for h in hurdles if h.get("gp_pct")), None)
            gp = (match.get("gp_pct") if match else "") or "0"
            gp = gp.replace("%", "")
            shares_by_name[name] = [
                gp if i < stage_n - 1 else final_gp for i in range(stage_n)
            ]

    return {
        "deal_name": deal,
        "deal_id": (classes[0].get("deal_id") if classes else "") or "",
        "mapped_from": path.name,
        "note": (
            "PUT into class-setup meta.promote (resolve share keys from class "
            "names → class UUIDs). Fixes Class Setup UI showing synthetic 12% "
            "instead of Class Mapping hurdle rates."
        ),
        "promote": {
            "hurdles": [
                {
                    "id": h["id"],
                    "rate": h["rate"],
                    "basis": h["basis"],
                    "measuredOn": h["measuredOn"],
                }
                for h in promote_hurdles
            ],
            "shares_by_class_name": shares_by_name,
        },
        "ui_maps_to": "add_deal_form.class_setup_json.promote",
    }


def write(payload: dict, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "ETL_Class_Promote_FILLED.payload.json"
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"deal={payload['deal_name']} hurdles={len(payload['promote']['hurdles'])}")
    print(f"wrote={path}")


def resolve_input(explicit: str | None) -> Path:
    if explicit:
        p = Path(explicit).expanduser().resolve()
        if not p.exists():
            raise FileNotFoundError(p)
        return p
    files = [f for f in INPUT_DIR.glob("Class_Mapping*.xlsx") if not f.name.startswith("~")]
    if not files:
        raise FileNotFoundError(f"No Class_Mapping_FILLED in {INPUT_DIR}")
    return max(files, key=lambda f: f.stat().st_mtime)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("class_mapping", nargs="?")
    parser.add_argument("--out-dir", default=str(OUTPUT_DIR))
    parser.add_argument("--all-migration-deals", action="store_true")
    args = parser.parse_args(argv)
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = Path(args.out_dir)

    if args.all_migration_deals:
        for deal, paths in MIGRATION_SOURCES.items():
            src = next((p for p in paths if p.exists()), None)
            if not src:
                print(f"SKIP {deal}")
                continue
            print(f"\n=== {deal} ===")
            payload = build_promote(src)
            write(payload, out / deal.lower().replace(" ", "_"))
        return 0

    src = resolve_input(args.class_mapping)
    write(build_promote(src), out)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        raise SystemExit(1)
