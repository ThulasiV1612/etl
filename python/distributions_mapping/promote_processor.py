"""
Class_Mapping_FILLED Hurdles / Final Tier → ETL_Class_Promote_FILLED.payload.json

Writes promote schedule under mapFolder/<deal_name>/ so Class Setup UI shows
real hurdle rates (not synthetic defaults).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
OUT_PAYLOAD = "ETL_Class_Promote_FILLED.payload.json"
OUT_WARNINGS = "warnings.txt"


def safe_folder_name(name: str) -> str:
    raw = str(name or "").strip() or "Unknown_Deal"
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", raw)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    cleaned = cleaned[:120] if cleaned else "Unknown_Deal"
    return cleaned or "Unknown_Deal"


def cell(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def read_sheet(path: Path, name: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if name not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    headers = [cell(h) for h in next(rows)]
    out: list[dict] = []
    for row in rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        if cell(row[0]).upper() in {"LOOKUP", "HINT"}:
            continue
        out.append(
            {
                headers[i]: cell(row[i]) if i < len(row) else ""
                for i in range(len(headers))
                if headers[i]
            }
        )
    wb.close()
    return out


def basis_from_pref_type(raw: str) -> str:
    t = (raw or "").lower()
    if "irr" in t:
        return "IRR"
    if "cash" in t or "coc" in t:
        return "Cash-on-cash"
    return "Cash-on-cash"


def build_promote(path: Path, deal_id: str | None = None) -> dict:
    classes = read_sheet(path, "Classes")
    hurdles = read_sheet(path, "Hurdles")
    final = read_sheet(path, "Final Tier")
    deal = (classes[0].get("deal_name") if classes else "") or path.stem
    resolved_deal_id = (
        str(deal_id or "").strip()
        or (classes[0].get("deal_id") if classes else "")
        or ""
    )

    promote_hurdles = []
    seen = set()
    for h in hurdles:
        if (h.get("is_final") or "").lower() in {"yes", "true", "1"}:
            continue
        rate = (h.get("hurdle_rate_pct") or "").replace("%", "").strip()
        if not rate or rate in seen:
            continue
        seen.add(rate)
        promote_hurdles.append(
            {
                "id": f"h{len(promote_hurdles) + 1}",
                "rate": rate,
                "basis": basis_from_pref_type(h.get("preferred_return_type") or ""),
                "measuredOn": "LP classes",
                "source_class_name": h.get("class_name") or "",
                "lp_pct": (h.get("lp_pct") or "").replace("%", ""),
                "gp_pct": (h.get("gp_pct") or "").replace("%", ""),
            }
        )

    ft = final[0] if final else {}
    final_lp = (ft.get("final_lp_pct") or "50").replace("%", "")
    final_gp = (ft.get("final_gp_pct") or "50").replace("%", "")

    shares_by_name: dict[str, list[str]] = {}
    stage_n = max(1, len(promote_hurdles)) + 1
    for c in classes:
        name = c.get("name") or ""
        if not name:
            continue
        st = (c.get("subscription_type") or "").lower()
        if st == "lp":
            match = next((h for h in hurdles if h.get("class_name") == name), None)
            lp = (match.get("lp_pct") if match else "") or "100"
            lp = lp.replace("%", "")
            shares_by_name[name] = [
                lp if i < stage_n - 1 else final_lp for i in range(stage_n)
            ]
        elif st == "gp":
            match = next((h for h in hurdles if h.get("gp_pct")), None)
            gp = (match.get("gp_pct") if match else "") or "0"
            gp = gp.replace("%", "")
            shares_by_name[name] = [
                gp if i < stage_n - 1 else final_gp for i in range(stage_n)
            ]

    return {
        "deal_name": deal,
        "deal_id": resolved_deal_id,
        "mapped_from": path.name,
        "note": (
            "PUT into class-setup meta.promote (resolve share keys from class "
            "names → class UUIDs). Fixes Class Setup UI showing synthetic rates "
            "instead of Class Mapping hurdle rates."
        ),
        "promote": {
            "hurdles": [
                {
                    "id": h["id"],
                    "rate": h["rate"],
                    "basis": h["basis"],
                    "measuredOn": h["measuredOn"],
                }
                for h in promote_hurdles
            ],
            "shares_by_class_name": shares_by_name,
        },
        "ui_maps_to": "add_deal_form.class_setup_json.promote",
    }


def process_workbook(
    input_path: Path,
    out_dir: Path,
    deal_id: str | None = None,
) -> dict:
    if not input_path.exists():
        raise FileNotFoundError(input_path)
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("Please upload an Excel (.xlsx) Class Mapping file.")

    classes = read_sheet(input_path, "Classes")
    if not classes:
        raise ValueError("No class rows found in the Class Mapping workbook.")

    payload = build_promote(input_path, deal_id=deal_id)
    map_root = Path(out_dir).expanduser().resolve()
    map_root.mkdir(parents=True, exist_ok=True)
    folder_name = safe_folder_name(payload["deal_name"])
    deal_dir = map_root / folder_name
    deal_dir.mkdir(parents=True, exist_ok=True)

    payload_path = deal_dir / OUT_PAYLOAD
    warnings_path = deal_dir / OUT_WARNINGS
    payload_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    warnings: list[str] = []
    if not payload.get("deal_id"):
        warnings.append("deal_id is empty — select a deal before loading promote into the database.")
    if not payload["promote"]["hurdles"]:
        warnings.append("No promote hurdles found (Hurdles sheet empty or only final rows).")
    if warnings:
        warnings_path.write_text("\n".join(warnings) + "\n", encoding="utf-8")
    elif warnings_path.exists():
        warnings_path.unlink()

    return {
        "success": True,
        "sourceFile": input_path.name,
        "summary": {
            "dealName": payload["deal_name"],
            "dealId": payload.get("deal_id") or "",
            "hurdles": len(payload["promote"]["hurdles"]),
            "shareClasses": len(payload["promote"]["shares_by_class_name"]),
            "warnings": len(warnings),
        },
        "warnings": warnings,
        "dealFolders": [
            {
                "dealName": payload["deal_name"],
                "folderName": folder_name,
                "files": {
                    "payload": str(payload_path),
                    "warnings": str(warnings_path) if warnings else None,
                },
            }
        ],
        "files": {
            "payload": str(payload_path),
            "warnings": str(warnings_path) if warnings else None,
        },
        "payload": payload,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Map Class_Mapping Hurdles/Final Tier → Class Setup promote payload"
    )
    parser.add_argument("class_mapping", help="Class_Mapping_FILLED.xlsx path")
    parser.add_argument("--out-dir", required=True, help="mapFolder root")
    parser.add_argument("--deal-id", default="", help="Override deal UUID in output")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    try:
        result = process_workbook(
            Path(args.class_mapping).expanduser().resolve(),
            Path(args.out_dir).expanduser().resolve(),
            deal_id=args.deal_id or None,
        )
    except Exception as e:
        message = str(e)
        print(json.dumps({"success": False, "error": message}), file=sys.stderr)
        print("ERROR:", message, file=sys.stderr)
        return 1

    if args.json:
        print(json.dumps(result))
    else:
        print("deal", result["summary"]["dealName"])
        print("hurdles", result["summary"]["hurdles"])
        print("wrote", result["files"]["payload"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
