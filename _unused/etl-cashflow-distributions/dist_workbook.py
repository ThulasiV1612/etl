"""Write ETL_Distribution_Setup_FILLED.xlsx + API payload helpers."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(color="FFFFFF", bold=True)
example_fill = PatternFill("solid", fgColor="E2EFDA")
hint_fill = PatternFill("solid", fgColor="FFF2CC")
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

WATERFALL_HEADERS = [
    "deal_name",
    "deal_id",
    "row_order",
    "row_id",
    "kind",
    "name",
    "pay_to_class_name",
    "pay_to_class_id",
    "amount_mode",
    "input_amount",
    "catchup_pct",
    "maps_to",
    "source_from_class_mapping",
]

KIND_HINTS = [
    "LOOKUP",
    "",
    "",
    "",
    "LP_PREF | PREF_CURRENT | PREF_ACCRUED | ROC | CATCHUP",
    "Display label",
    "Exact Class Setup name (resolved to UUID on load)",
    "Fill UUID on load OR leave blank for API resolve",
    "calc | input",
    "Used when amount_mode=input",
    "CATCHUP only (default 20)",
    "distribution_setup_json.waterfalls.*",
    "Do not import yellow hint rows",
]


def style_header(ws, cols: int) -> None:
    ws.row_dimensions[1].height = 36
    for c in range(1, cols + 1):
        cell = ws.cell(1, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def mark(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = example_fill
        cell.border = thin


def hint_row(ws, values: list) -> None:
    ws.append(values)
    for c in range(1, len(values) + 1):
        cell = ws.cell(ws.max_row, c)
        cell.fill = hint_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = thin
    ws.row_dimensions[ws.max_row].height = 40


def autosize(ws, max_w: int = 34) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        w = 12
        for cell in col:
            if cell.value is not None:
                w = max(w, min(max_w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = w


def slug(name: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "_", name).strip("_").lower()
    return s or "class"


@dataclass
class WaterfallRow:
    deal_name: str
    deal_id: str
    row_order: int
    row_id: str
    kind: str
    name: str
    pay_to_class_name: str
    amount_mode: str = "calc"
    input_amount: str = "0"
    catchup_pct: str = ""
    source: str = ""
    pay_to_class_id: str = ""

    def as_list(self, maps_to: str) -> list:
        return [
            self.deal_name,
            self.deal_id,
            self.row_order,
            self.row_id,
            self.kind,
            self.name,
            self.pay_to_class_name,
            self.pay_to_class_id,
            self.amount_mode,
            self.input_amount,
            self.catchup_pct,
            maps_to,
            self.source,
        ]

    def api_row(self) -> dict:
        names = [x.strip() for x in str(self.pay_to_class_name).split(";") if x.strip()]
        out = {
            "id": self.row_id,
            "kind": self.kind,
            "name": self.name,
            "payTo": names,
            "amountMode": self.amount_mode,
            "inputAmount": self.input_amount,
        }
        if self.kind == "CATCHUP":
            out["catchupPct"] = self.catchup_pct or "20"
        return out


@dataclass
class DistSetupMapped:
    deal_name: str
    deal_id: str = ""
    target_raise: str = "0"
    mapped_from: str = "Class_Mapping_FILLED.xlsx"
    operating: list[WaterfallRow] = field(default_factory=list)
    capital: list[WaterfallRow] = field(default_factory=list)
    class_setup_context: dict = field(default_factory=dict)

    def payload(self) -> dict:
        return {
            "deal_name": self.deal_name,
            "deal_id": self.deal_id,
            "target_raise": self.target_raise,
            "mapped_from": self.mapped_from,
            "note": "payTo holds class NAMES — replace with deal_investor_class.id UUIDs before PUT",
            "class_setup_context": self.class_setup_context,
            "waterfalls": {
                "operating": [r.api_row() for r in self.operating],
                "capital": [r.api_row() for r in self.capital],
            },
        }


def write_blank_template(path: Path) -> None:
    """Blank ETL Distribution Setup template (yellow hints = docs)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    ws["A1"] = "ETL: Class_Mapping_FILLED → Syndication Distribution Setup"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    lines = [
        "",
        "SOURCE OF TRUTH: Class_Mapping_FILLED.xlsx (not Cash Flow paid Excel).",
        "DB TARGET: add_deal_form.distribution_setup_json.waterfalls",
        "API: PUT /deals/:dealId/distribution-setup  { waterfalls: { operating, capital } }",
        "",
        "LOAD ORDER:",
        "1. Deal exists + Classes loaded (names must match pay_to_class_name)",
        "2. Fill / generate Waterfall_Operating + Waterfall_Capital",
        "3. Resolve pay_to_class_name → class UUID → payTo[]",
        "4. PUT distribution-setup",
        "",
        "Yellow rows = documentation — do not import.",
        "Green rows = data to load.",
        "Promote / final split stay on Class Setup — not in this workbook.",
        "Cash Flow distributions paid Excel = ledger only (optional priorDistributions later).",
    ]
    for i, line in enumerate(lines, start=2):
        ws[f"A{i}"] = line
    ws.column_dimensions["A"].width = 100

    ws = wb.create_sheet("Deal_Meta")
    ws.append(["deal_name", "deal_id", "target_raise", "mapped_from", "notes"])
    style_header(ws, 5)
    hint_row(
        ws,
        [
            "LOOKUP",
            "DB deal UUID",
            "From Classes.deal_target_raise",
            "Class_Mapping_FILLED.xlsx",
            "Fill deal_id before API load",
        ],
    )
    autosize(ws)

    for title, maps_to in (
        ("Waterfall_Operating", "distribution_setup_json.waterfalls.operating[]"),
        ("Waterfall_Capital", "distribution_setup_json.waterfalls.capital[]"),
    ):
        ws = wb.create_sheet(title)
        ws.append(WATERFALL_HEADERS)
        style_header(ws, len(WATERFALL_HEADERS))
        hint = list(KIND_HINTS)
        hint[11] = maps_to
        hint_row(ws, hint)
        autosize(ws)

    ws = wb.create_sheet("Kind_Map")
    ws.append(["cash_flow_or_class_idea", "our_kind", "waterfall", "notes"])
    style_header(ws, 4)
    for row in [
        ["LP preferred return / CoC pref", "LP_PREF", "operating", "From Classes.pref_enabled"],
        ["Preferred equity current coupon", "PREF_CURRENT", "operating", "preferred_equity class"],
        ["Preferred equity accrued", "PREF_ACCRUED", "capital", "preferred_equity class"],
        ["Return of capital", "ROC", "capital", "Each LP class"],
        ["Carried interest / GP catch-up", "CATCHUP", "capital", "GP + Final Tier %"],
        ["Acquisition fee / Other payment", "(ledger only)", "—", "Not a setup kind"],
    ]:
        ws.append(row)
        mark(ws, ws.max_row, 4)
    autosize(ws)

    ws = wb.create_sheet("Load_Steps")
    ws.append(["step", "action"])
    style_header(ws, 2)
    for row in [
        ["1", "Ensure deal classes loaded from Class_Mapping_FILLED"],
        ["2", "Generate this FILLED file via map_class_to_distribution_setup.py"],
        ["3", "Resolve pay_to_class_name → UUID (load_distribution_setup.py)"],
        ["4", "PUT /deals/:dealId/distribution-setup"],
        ["5", "Confirm Deal → Distribution Setup UI"],
    ]:
        ws.append(row)
        mark(ws, ws.max_row, 2)
    autosize(ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_filled_workbook(mapped: DistSetupMapped, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Deal_Meta"
    ws.append(["deal_name", "deal_id", "target_raise", "mapped_from", "notes"])
    style_header(ws, 5)
    ws.append(
        [
            mapped.deal_name,
            mapped.deal_id,
            mapped.target_raise,
            mapped.mapped_from,
            "Syndication Distribution Setup seed from Class Mapping. Promote stays on Class Setup.",
        ]
    )
    mark(ws, 2, 5)
    autosize(ws)

    for title, rows, maps_to in (
        (
            "Waterfall_Operating",
            mapped.operating,
            "distribution_setup_json.waterfalls.operating[]",
        ),
        (
            "Waterfall_Capital",
            mapped.capital,
            "distribution_setup_json.waterfalls.capital[]",
        ),
    ):
        ws = wb.create_sheet(title)
        ws.append(WATERFALL_HEADERS)
        style_header(ws, len(WATERFALL_HEADERS))
        hint = list(KIND_HINTS)
        hint[11] = maps_to
        hint_row(ws, hint)
        for i, r in enumerate(rows, start=3):
            ws.append(r.as_list(maps_to))
            mark(ws, i, len(WATERFALL_HEADERS))
        autosize(ws)

    ws = wb.create_sheet("Load_Steps")
    ws.append(["step", "action"])
    style_header(ws, 2)
    for row in [
        ["1", f"Ensure {mapped.deal_name} classes loaded from Class_Mapping_FILLED"],
        ["2", "Fill deal_id + resolve pay_to_class_id UUIDs (or use load_distribution_setup.py)"],
        ["3", "PUT /deals/:dealId/distribution-setup with payload JSON"],
        ["4", "Confirm Deal → Distribution Setup UI"],
        ["5", "Promote/final split already on Class Setup — not in distribution_setup_json"],
    ]:
        ws.append(row)
        mark(ws, ws.max_row, 2)
    autosize(ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_payload(mapped: DistSetupMapped, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(mapped.payload(), indent=2), encoding="utf-8")
