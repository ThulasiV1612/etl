"""
Cash Flow 'distributions paid' Excel → ETL_Distributions_Paid_FILLED.xlsx + .payload.json

Writes under mapFolder/<deal_name>/ (same pattern as Distribution Setup mapping).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent

OUT_XLSX = "ETL_Distributions_Paid_FILLED.xlsx"
OUT_PAYLOAD = "ETL_Distributions_Paid_FILLED.payload.json"
OUT_WARNINGS = "warnings.txt"

SHEET_CANDIDATES = (
    "Distributions by investment",
    "Paid_By_Investment",
)

# Cash Flow single-period export sheets start with "Payment details"
PAYMENT_DETAILS_PREFIX = "payment details"

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(color="FFFFFF", bold=True)
example_fill = PatternFill("solid", fgColor="E2EFDA")
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def safe_folder_name(name: str) -> str:
    raw = str(name or "").strip() or "Unknown_Deal"
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", raw)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    cleaned = cleaned[:120] if cleaned else "Unknown_Deal"
    return cleaned or "Unknown_Deal"


def style_header(ws, cols: int) -> None:
    ws.row_dimensions[1].height = 36
    for c in range(1, cols + 1):
        cell = ws.cell(1, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def mark(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = example_fill
        cell.border = thin


def autosize(ws, max_w: int = 32) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        w = 12
        for cell in col[:120]:
            if cell.value is not None:
                w = max(w, min(max_w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = w


def cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if s in {"-", "—", "None", "null", "n/a", "N/A"}:
        return ""
    return s


def money(v) -> float:
    s = cell(v).replace("$", "").replace(",", "").strip()
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def to_date(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = cell(v)
    if not s:
        return ""
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        return s[:10]
    for fmt in ("%m/%d/%Y", "%m-%d-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return s[:10]


def map_distribution_type(dist_type: str, source: str) -> str:
    t = (dist_type or "").strip().lower()
    if source in {"fee", "distribution_fee"}:
        return "acquisition_fee"
    if not t:
        if source == "gp_payment":
            return "return_on_capital"
        return "return_of_capital" if source == "capital" else "preferred_return"
    if "acquisition" in t and "fee" in t:
        return "acquisition_fee"
    if re.search(r"return\s+on\s+capital", t):
        return "return_on_capital"
    if "return of capital" in t or t == "roc":
        return "return_of_capital"
    if "promote" in t or "carried" in t:
        return "promote"
    if "preferred" in t or t == "pref":
        return "preferred_return"
    if "other" in t:
        return "other"
    return (dist_type or "").strip() or "preferred_return"


def map_deducts_from(distribution_type: str, source: str) -> str:
    t = (distribution_type or "").strip().lower()
    src = (source or "").strip().lower()
    if src in {"fee", "distribution_fee"} or "acquisition" in t or t == "acquisition_fee":
        return "fee"
    if (
        src == "capital"
        or t == "return_of_capital"
        or "return of capital" in t
        or t == "roc"
    ):
        return "capital"
    if src in {"gp_payment", "gp payment"}:
        return "accrued_pref"
    if "current" in t:
        return "current_pref"
    return "accrued_pref"


DISTRIBUTE_SUMMARY_RE = re.compile(
    r"^Distribute\s+100%\s+to\s+(.+?)\s+-\s+summary\s*$",
    flags=re.I,
)


def is_acquisition_fee(dist_type: str, memo: str = "") -> bool:
    """
    Cash Flow GP acquisition payouts → portal source=fee + distributionFee.
    Matches memo/type like 'GP Acquisition Fee', 'Harris Ridge Acquisition Fee',
    'Cottonwood Acquisition Share' (not only names that contain 'fee').
    """
    blob = f"{dist_type or ''} {memo or ''}".strip().lower()
    if not blob:
        return False
    if "acquisition" in blob and ("fee" in blob or "share" in blob):
        return True
    # Explicit acquisition fee type from Cash Flow
    if "acquisition" in (dist_type or "").strip().lower() and "fee" in (
        dist_type or ""
    ).strip().lower():
        return True
    return False


def is_gp_class_name(name: str) -> bool:
    n = (name or "").strip().lower()
    if not n:
        return False
    return (
        "general partner" in n
        or n in {"gp", "class b - gp", "class b gp"}
        or (re.search(r"\bgp\b", n) is not None and "lp" not in n)
    )


def unique_class_names(names: list[str]) -> list[str]:
    out: list[str] = []
    for name in names:
        n = str(name or "").strip()
        if n and n not in out:
            out.append(n)
    return out


def find_distribute_100_class_names(row: dict) -> list[str]:
    """Headers like 'Distribute 100% to Class B - General partners - summary'."""
    found: list[str] = []
    for key, value in (row or {}).items():
        m = DISTRIBUTE_SUMMARY_RE.match(str(key or "").strip())
        if not m:
            continue
        summary = cell(value).lower()
        # Cash Flow uses Split / 100% / similar when this hurdle applies
        if summary in {"", "-", "—", "0", "none", "n/a"}:
            continue
        class_name = m.group(1).strip()
        if class_name and class_name not in found:
            found.append(class_name)
    return found


def build_distribution_fee_config(
    *,
    dist_type: str,
    memo: str,
    period_start: str,
    period_end: str,
    amount: float,
    class_names: list[str],
) -> dict | None:
    if not class_names:
        return None
    if not is_acquisition_fee(dist_type, memo):
        return None
    # 100% of fee cash to the fee class(es); equal split if multiple.
    pct = round(100.0 / len(class_names), 6) if class_names else 0
    splits = []
    remaining = 100.0
    for i, name in enumerate(class_names):
        share = remaining if i == len(class_names) - 1 else pct
        remaining = round(remaining - share, 6)
        splits.append(
            {
                "classId": "",
                "className": name,
                "percent": str(share).rstrip("0").rstrip(".") if "." in str(share) else str(int(share)),
            }
        )
    fee_name = memo or dist_type or "Distribution fee"
    fee_type = "Acquisition fee"
    return {
        "name": fee_name,
        "type": fee_type,
        "typeOptions": [fee_type],
        "cashAvailable": str(round(amount, 2)),
        "periodFactor": "0.083333",
        "periodStart": period_start or "",
        "periodEnd": period_end or "",
        "classSplits": splits,
        "maps_to": "distribution_setup_json.distributionFee (Distribution Fee tab)",
    }


def infer_source(
    dist_type: str,
    pref_paid: float,
    roc: float,
    carried: float,
    other: float,
    *,
    class_names: list[str] | None = None,
    memo: str = "",
) -> str:
    t = (dist_type or "").strip().lower()
    # GP Acquisition Fee → Distribution Fee tab (source=fee), not operating / gp_payment
    if is_acquisition_fee(dist_type, memo):
        return "fee"
    if "return of capital" in t or t == "roc" or (roc > 0 and pref_paid <= 0):
        return "capital"
    if "carried" in t or (carried > 0 and pref_paid <= 0 and other <= 0):
        return "capital"
    if "acquisition" in t or "other" in t:
        return "operating"
    if pref_paid > 0 or "preferred" in t:
        return "operating"
    if roc > 0:
        return "capital"
    return "operating"


def find_header_row(ws, max_scan: int = 25) -> tuple[int, list[str]]:
    """Return 0-based header row index and header labels (Cash Flow may put headers mid-sheet)."""
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True)):
        headers_raw = [cell(h) for h in (row or [])]
        lowered = [h.lower() for h in headers_raw]
        has_deal = any(h == "deal name" for h in lowered)
        has_uid = any(
            h in {"unique investment id", "source investment id"} for h in lowered
        )
        has_net = any(h == "net payment" for h in lowered)
        has_memo = any(h == "distribution memo" for h in lowered)
        if has_deal and has_uid and (has_net or has_memo):
            headers = []
            for h in headers_raw:
                m = re.match(r"^(.+?)\s*\[(.+)\]\s*$", h)
                headers.append(m.group(2) if m else h)
            return idx, headers
    raise ValueError(
        "Could not find Cash Flow investor header row "
        "(expected Deal name + Unique investment ID + Net payment)."
    )


def pick_cashflow_sheet(wb) -> str:
    sheet_name = next((n for n in SHEET_CANDIDATES if n in wb.sheetnames), None)
    if sheet_name:
        return sheet_name
    for n in wb.sheetnames:
        if n.strip().lower().startswith(PAYMENT_DETAILS_PREFIX):
            return n
    raise ValueError(
        "No 'Distributions by investment' or 'Payment details…' sheet found. "
        "Upload a Cash Flow distributions-paid or single-period distribution Excel."
    )


def read_distributions_by_investment(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet_name = pick_cashflow_sheet(wb)
        ws = wb[sheet_name]
        header_idx, headers = find_header_row(ws)
        out: list[dict] = []
        for row in ws.iter_rows(min_row=header_idx + 2, values_only=True):
            if not row or all(c is None or str(c).strip() in ("", "-") for c in row[:4]):
                continue
            out.append(
                {
                    headers[i]: row[i]
                    for i in range(min(len(headers), len(row)))
                    if headers[i]
                }
            )
        return out
    finally:
        wb.close()


def build_runs(rows: list[dict], *, source_file: str, deal_id: str = "") -> dict:
    groups: dict[tuple, list[dict]] = defaultdict(list)
    deal_name = ""

    for r in rows:
        deal = cell(r.get("Deal name"))
        if deal:
            deal_name = deal
        memo = cell(r.get("Distribution memo")) or cell(r.get("Distribution type")) or "Distribution"
        dist_type = cell(r.get("Distribution type"))
        period_start = to_date(r.get("Period start date"))
        period_end = to_date(r.get("Period end date") or r.get("Payment initiated on"))
        payment_date = to_date(r.get("Payment initiated on") or r.get("Period end date"))
        net = money(r.get("Net payment"))
        if net < 0:
            continue
        # Keep $0 lines so every investor on the run appears in Distribution Details
        has_investor = bool(
            cell(r.get("Unique investment ID"))
            or cell(r.get("Investor name"))
            or cell(r.get("Investment profile name"))
        )
        if not has_investor:
            continue
        groups[(period_end, memo, dist_type, period_start)].append(
            {**r, "_period_start": period_start, "_period_end": period_end, "_payment_date": payment_date}
        )

    runs: list[dict] = []
    for (period_end, memo, dist_type, period_start), lines in sorted(
        groups.items(), key=lambda kv: (kv[0][0], kv[0][1])
    ):
        investor_lines = []
        total = 0.0
        pref_sum = roc_sum = carried_sum = other_sum = 0.0
        payment_date = ""
        for r in lines:
            if not payment_date:
                payment_date = cell(r.get("_payment_date"))
            net = money(r.get("Net payment"))
            pref = money(r.get("Preferred return paid (by distribution)"))
            roc = money(r.get("Return of capital"))
            carried = money(
                r.get("Carried interest paid") if cell(r.get("Carried interest paid")) else 0
            )
            other = money(r.get("Other payment"))
            pref_sum += pref
            roc_sum += roc
            carried_sum += carried if carried > 0 else 0
            other_sum += other
            total += net
            capital = money(r.get("Outstanding capital") or r.get("Initial investment amount"))
            pct = money(r.get("Percentage of offering"))
            pct_display = round(pct * 100, 6) if pct <= 1 else round(pct, 6)
            investor_lines.append(
                {
                    "investorId": "",
                    "contactId": "",
                    "userEmail": "",
                    "investorName": cell(r.get("Investor name")),
                    "classId": "",
                    "className": cell(r.get("Class name")),
                    "capital": str(round(capital, 2)),
                    "percentOfClass": str(pct_display),
                    "payment": str(round(net, 2)),
                    "source_investment_id": cell(r.get("Unique investment ID")),
                    "profileName": cell(r.get("Investment profile name")),
                    "distribution_method": cell(
                        r.get("Distribution method") or r.get("Selected distribution method")
                    ),
                }
            )

        if total <= 0:
            continue

        fee_class_names: list[str] = []
        for r in lines:
            for name in find_distribute_100_class_names(r):
                if name not in fee_class_names:
                    fee_class_names.append(name)

        class_names = unique_class_names(
            [
                str(line.get("className") or "").strip()
                for line in investor_lines
                if str(line.get("className") or "").strip()
            ]
        )
        # Acquisition fee (type or memo like "GP Acquisition Fee") → source=fee + distributionFee.
        # Fallback class = paid class(es) when Cash Flow has no "Distribute 100%" columns.
        is_fee = is_acquisition_fee(dist_type, memo)
        if is_fee and not fee_class_names:
            fee_class_names = list(class_names)
        source = (
            "fee"
            if is_fee and fee_class_names
            else infer_source(
                dist_type,
                pref_sum,
                roc_sum,
                carried_sum,
                other_sum,
                class_names=class_names,
                memo=memo,
            )
        )
        run_date = period_end or payment_date or datetime.now().strftime("%Y-%m-%d")
        portal_type = map_distribution_type(dist_type, source)
        fee_config = (
            build_distribution_fee_config(
                dist_type=dist_type,
                memo=memo,
                period_start=period_start,
                period_end=period_end or run_date,
                amount=total,
                class_names=fee_class_names,
            )
            if source == "fee"
            else None
        )
        runs.append(
            {
                "source": source,
                "amount": str(round(total, 2)),
                "date": run_date,
                "name": memo,
                "notes": (
                    f"Cash Flow import | type={dist_type or '—'} | "
                    f"pref={round(pref_sum, 2)} roc={round(roc_sum, 2)} "
                    f"other={round(other_sum, 2)} carried={round(carried_sum, 2)} | "
                    f"file={source_file}"
                    + (
                        f" | distributionFee classes={', '.join(fee_class_names)}"
                        if fee_config
                        else ""
                    )
                ),
                "period": "monthly",
                "periodStart": period_start,
                "periodEnd": period_end or run_date,
                "paymentDate": payment_date or run_date,
                "distributionType": portal_type,
                "deductsFrom": map_deducts_from(portal_type, source),
                "visible": True,
                "cash_flow_distribution_type": dist_type,
                "investorPayments": investor_lines,
                "line_count": len(investor_lines),
                "is_distribution_fee": bool(fee_config),
                "distribution_fee": fee_config,
            }
        )

    fee_cfgs = [r.get("distribution_fee") for r in runs if r.get("distribution_fee")]
    return {
        "deal_name": deal_name.strip(),
        "deal_id": str(deal_id or "").strip(),
        "mapped_from": source_file,
        "note": (
            "Historical distributions already paid. Load via Upload Hub → Upload Prior Distributions. "
            "Acquisition fee + 'Distribute 100% to … - summary' → distribution_setup_json.distributionFee. "
            "Does NOT replace waterfall Distribution Setup."
        ),
        "prior_distributions": runs,
        "distribution_fee": fee_cfgs[-1] if fee_cfgs else None,
        "run_count": len(runs),
        "total_paid": round(sum(float(r["amount"]) for r in runs), 2),
    }


def write_filled_xlsx(payload: dict, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Deal_Meta"
    ws.append(["deal_name", "deal_id", "source_file", "run_count", "total_paid", "notes"])
    style_header(ws, 6)
    ws.append(
        [
            payload["deal_name"],
            payload.get("deal_id", ""),
            payload["mapped_from"],
            payload["run_count"],
            payload["total_paid"],
            "Already-paid ledger → priorDistributions (not waterfall setup)",
        ]
    )
    mark(ws, 2, 6)
    autosize(ws)

    run_headers = [
        "run_order",
        "date",
        "name",
        "source",
        "amount",
        "period_start",
        "period_end",
        "payment_date",
        "distribution_type",
        "deducts_from",
        "visible",
        "cash_flow_distribution_type",
        "line_count",
        "period",
        "notes",
        "maps_to",
    ]
    ws = wb.create_sheet("Prior_Runs")
    ws.append(run_headers)
    style_header(ws, len(run_headers))
    for i, r in enumerate(payload["prior_distributions"], start=1):
        ws.append(
            [
                i,
                r["date"],
                r["name"],
                r["source"],
                r["amount"],
                r.get("periodStart", ""),
                r.get("periodEnd", ""),
                r.get("paymentDate", ""),
                r.get("distributionType", ""),
                r.get("deductsFrom", ""),
                "yes" if r.get("visible", True) else "no",
                r.get("cash_flow_distribution_type", ""),
                r["line_count"],
                r.get("period", ""),
                r.get("notes", ""),
                "Deal → Distributions table (priorDistributions[])",
            ]
        )
        mark(ws, ws.max_row, len(run_headers))
    autosize(ws)

    line_headers = [
        "run_order",
        "run_date",
        "run_name",
        "run_amount",
        "source_investment_id",
        "investor_name",
        "profile_name",
        "class_name",
        "class_id",
        "investor_id",
        "capital",
        "percent_of_class",
        "payment",
        "distribution_method",
        "maps_to",
    ]
    ws = wb.create_sheet("Prior_Investor_Lines")
    ws.append(line_headers)
    style_header(ws, len(line_headers))
    for i, r in enumerate(payload["prior_distributions"], start=1):
        for line in r["investorPayments"]:
            ws.append(
                [
                    i,
                    r["date"],
                    r["name"],
                    r["amount"],
                    line.get("source_investment_id", ""),
                    line.get("investorName", ""),
                    line.get("profileName", ""),
                    line.get("className", ""),
                    line.get("classId", ""),
                    line.get("investorId", ""),
                    line.get("capital", 0),
                    line.get("percentOfClass", 0),
                    line.get("payment", 0),
                    line.get("distribution_method", ""),
                    "priorDistributions[].investorPayments[]",
                ]
            )
            mark(ws, ws.max_row, len(line_headers))
    autosize(ws)

    fee = payload.get("distribution_fee")
    if fee:
        ws = wb.create_sheet("Distribution_Fee")
        fee_headers = [
            "name",
            "type",
            "cash_available",
            "period_factor",
            "period_start",
            "period_end",
            "class_name",
            "class_id",
            "percent",
            "maps_to",
        ]
        ws.append(fee_headers)
        style_header(ws, len(fee_headers))
        splits = fee.get("classSplits") or []
        if not splits:
            splits = [{"className": "", "classId": "", "percent": "100"}]
        for i, split in enumerate(splits):
            ws.append(
                [
                    fee.get("name", "") if i == 0 else "",
                    fee.get("type", "") if i == 0 else "",
                    fee.get("cashAvailable", "") if i == 0 else "",
                    fee.get("periodFactor", "") if i == 0 else "",
                    fee.get("periodStart", "") if i == 0 else "",
                    fee.get("periodEnd", "") if i == 0 else "",
                    split.get("className", ""),
                    split.get("classId", ""),
                    split.get("percent", ""),
                    "distribution_setup_json.distributionFee (% Distribution Fee tab)",
                ]
            )
            mark(ws, ws.max_row, len(fee_headers))
        autosize(ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def parse_period_from_filename(name: str) -> tuple[str, str]:
    """Parse '… distribution (MM-DD-YYYY - MM-DD-YYYY)_….xlsx' → (start, end) ISO."""
    m = re.search(
        r"distribution\s*\(\s*(\d{1,2})[-/](\d{1,2})[-/](\d{4})\s*-\s*(\d{1,2})[-/](\d{1,2})[-/](\d{4})\s*\)",
        name or "",
        flags=re.I,
    )
    if not m:
        return "", ""
    m1, d1, y1, m2, d2, y2 = m.groups()
    return (
        f"{y1}-{int(m1):02d}-{int(d1):02d}",
        f"{y2}-{int(m2):02d}-{int(d2):02d}",
    )


def apply_filename_period(runs: list[dict], source_file: str) -> list[dict]:
    start, end = parse_period_from_filename(source_file)
    if not start and not end:
        return runs
    out = []
    for r in runs:
        row = dict(r)
        if not row.get("periodStart") and start:
            row["periodStart"] = start
        if not row.get("periodEnd") and end:
            row["periodEnd"] = end
        if not row.get("paymentDate") and end:
            row["paymentDate"] = end
        if not row.get("date") and end:
            row["date"] = end
        row["source_file"] = source_file
        out.append(row)
    return out


def map_one_workbook(input_path: Path, deal_id: str = "") -> dict:
    """Map a single Cash Flow paid / Payment-details Excel → prior run payload fragment."""
    if not input_path.exists():
        raise FileNotFoundError(input_path)
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError(f"Skip non-xlsx: {input_path.name}")

    rows = read_distributions_by_investment(input_path)
    if not rows:
        raise ValueError(f"No paid distribution rows found in {input_path.name}")

    payload = build_runs(
        rows,
        source_file=input_path.name,
        deal_id=deal_id or "",
    )
    if not payload["deal_name"]:
        payload["deal_name"] = input_path.stem.split("distribut")[0].strip(" _-") or "Deal"

    payload["prior_distributions"] = apply_filename_period(
        payload["prior_distributions"],
        input_path.name,
    )
    if not payload["prior_distributions"]:
        raise RuntimeError(
            f"No prior distribution runs in {input_path.name} (all Net payment ≤ 0)."
        )
    return payload


def merge_payloads(parts: list[dict], deal_id: str = "") -> dict:
    """Merge N mapped workbooks into one priorDistributions payload (any count)."""
    if not parts:
        raise ValueError("No distribution Excel files to map.")

    deal_name = next((p.get("deal_name") or "" for p in parts if p.get("deal_name")), "")
    deal = str(deal_id or "").strip() or next(
        (p.get("deal_id") or "" for p in parts if p.get("deal_id")), ""
    )
    sources = [p.get("mapped_from") or "" for p in parts if p.get("mapped_from")]
    runs: list[dict] = []
    seen: set[tuple] = set()

    for part in parts:
        for r in part.get("prior_distributions") or []:
            key = (
                str(r.get("periodStart") or ""),
                str(r.get("periodEnd") or ""),
                str(r.get("name") or ""),
                str(r.get("distributionType") or ""),
                str(r.get("amount") or ""),
            )
            if key in seen:
                continue
            seen.add(key)
            runs.append(r)

    runs.sort(key=lambda r: (str(r.get("periodEnd") or r.get("date") or ""), str(r.get("name") or "")))

    fee_cfgs = [
        p.get("distribution_fee")
        for p in parts
        if p.get("distribution_fee")
    ]
    # Prefer fee config from the latest period run inside merged runs
    for r in reversed(runs):
        if r.get("distribution_fee"):
            fee_cfgs.append(r["distribution_fee"])
            break

    return {
        "deal_name": deal_name.strip(),
        "deal_id": deal,
        "mapped_from": " | ".join(sources),
        "source_files": sources,
        "note": (
            "Each Cash Flow period / Payment-details Excel → one Deal → Distributions row. "
            "Acquisition fee + Distribute 100% columns → Distribution Fee tab "
            "(distribution_setup_json.distributionFee). "
            "investorPayments[] fills Distribution Details."
        ),
        "prior_distributions": runs,
        "distribution_fee": fee_cfgs[-1] if fee_cfgs else None,
        "run_count": len(runs),
        "total_paid": round(sum(float(r.get("amount") or 0) for r in runs), 2),
    }


def write_payload_outputs(payload: dict, out_dir: Path) -> dict:
    map_root = Path(out_dir).expanduser().resolve()
    map_root.mkdir(parents=True, exist_ok=True)
    folder_name = safe_folder_name(payload["deal_name"] or "Deal")
    deal_dir = map_root / folder_name
    deal_dir.mkdir(parents=True, exist_ok=True)

    xlsx_path = deal_dir / OUT_XLSX
    payload_path = deal_dir / OUT_PAYLOAD
    warnings_path = deal_dir / OUT_WARNINGS

    write_filled_xlsx(payload, xlsx_path)
    payload_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    warnings: list[str] = []
    if not payload.get("deal_id"):
        warnings.append("deal_id is empty — select a deal before loading to the database.")
    if payload["run_count"] == 0:
        warnings.append("No distribution runs mapped.")
    if warnings:
        warnings_path.write_text("\n".join(warnings) + "\n", encoding="utf-8")
    elif warnings_path.exists():
        warnings_path.unlink()

    return {
        "success": True,
        "sourceFile": payload.get("mapped_from") or "",
        "sourceFiles": payload.get("source_files") or [],
        "summary": {
            "dealName": payload["deal_name"],
            "dealId": payload.get("deal_id") or "",
            "runCount": payload["run_count"],
            "totalPaid": payload["total_paid"],
            "lineCount": sum(r.get("line_count", 0) for r in payload["prior_distributions"]),
            "fileCount": len(payload.get("source_files") or []),
            "warnings": len(warnings),
        },
        "warnings": warnings,
        "dealFolders": [
            {
                "dealName": payload["deal_name"],
                "folderName": folder_name,
                "files": {
                    "workbook": str(xlsx_path),
                    "payload": str(payload_path),
                    "warnings": str(warnings_path) if warnings else None,
                },
            }
        ],
        "files": {
            "workbook": str(xlsx_path),
            "payload": str(payload_path),
            "warnings": str(warnings_path) if warnings else None,
        },
    }


def collect_input_excels(paths: list[Path], input_dir: Path | None = None) -> list[Path]:
    files: list[Path] = []
    for p in paths:
        if p.is_file() and p.suffix.lower() == ".xlsx":
            files.append(p)
    if input_dir and input_dir.is_dir():
        for p in sorted(input_dir.glob("*.xlsx")):
            # Prefer period / payment-details / distributions-paid workbooks
            name = p.name.lower()
            if name.startswith("~$"):
                continue
            files.append(p)
    # de-dupe by resolved path
    seen: set[str] = set()
    out: list[Path] = []
    for p in files:
        key = str(p.resolve())
        if key in seen:
            continue
        seen.add(key)
        out.append(p.resolve())
    return out


def process_workbook(
    input_path: Path,
    out_dir: Path,
    deal_id: str | None = None,
) -> dict:
    return process_workbooks([input_path], out_dir, deal_id=deal_id)


def process_workbooks(
    input_paths: list[Path],
    out_dir: Path,
    deal_id: str | None = None,
    input_dir: Path | None = None,
) -> dict:
    files = collect_input_excels(input_paths, input_dir)
    if not files:
        raise ValueError(
            "No Excel (.xlsx) distribution files found. "
            "Upload one or more Cash Flow period files "
            "(e.g. 'Deal distribution (MM-DD-YYYY - MM-DD-YYYY)_….xlsx') "
            "or a distributions-paid workbook."
        )

    parts: list[dict] = []
    file_errors: list[str] = []
    for path in files:
        try:
            parts.append(map_one_workbook(path, deal_id=deal_id or ""))
        except Exception as e:
            file_errors.append(f"{path.name}: {e}")

    if not parts:
        raise RuntimeError(
            "No distribution runs could be mapped. "
            + (" | ".join(file_errors) if file_errors else "Check Net payment values.")
        )

    payload = merge_payloads(parts, deal_id=deal_id or "")
    result = write_payload_outputs(payload, out_dir)
    if file_errors:
        result["warnings"] = list(result.get("warnings") or []) + file_errors
        result["summary"]["warnings"] = len(result["warnings"])
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Map one or many Cash Flow distribution Excels → Prior Distributions FILLED files. "
            "Each period / Payment-details file becomes one Deal → Distributions row."
        )
    )
    parser.add_argument(
        "paid_excel",
        nargs="*",
        help="One or more Cash Flow distribution .xlsx files",
    )
    parser.add_argument(
        "--input-dir",
        default="",
        help="Optional folder of .xlsx files (e.g. Distributions/)",
    )
    parser.add_argument("--out-dir", required=True, help="mapFolder root")
    parser.add_argument("--deal-id", default="", help="Override deal UUID in output")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    try:
        result = process_workbooks(
            [Path(p).expanduser().resolve() for p in (args.paid_excel or [])],
            Path(args.out_dir).expanduser().resolve(),
            deal_id=args.deal_id or None,
            input_dir=Path(args.input_dir).expanduser().resolve() if args.input_dir else None,
        )
    except Exception as e:
        message = str(e)
        print(json.dumps({"success": False, "error": message}), file=sys.stderr)
        print("ERROR:", message, file=sys.stderr)
        return 1

    if args.json:
        print(json.dumps(result))
    else:
        print("deal", result["summary"]["dealName"])
        print("runs", result["summary"]["runCount"])
        print("files", result["summary"].get("fileCount", 1))
        print("total_paid", result["summary"]["totalPaid"])
        print("wrote", result["files"]["workbook"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
