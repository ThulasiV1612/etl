"""Write Class_Mapping_FILLED.xlsx (Classes / Hurdles / Final Tier)."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(color="FFFFFF", bold=True)
example_fill = PatternFill("solid", fgColor="E2EFDA")
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

CLASS_HEADERS = [
    "deal_name",
    "deal_id",
    "source_class_name",
    "name",
    "subscription_type",
    "status",
    "display_order",
    "entity_name",
    "start_date",
    "offering_size",
    "raise_amount_distributions",
    "billing_raise_quota",
    "minimum_investment",
    "number_of_units",
    "price_per_unit",
    "visibility",
    "equity_pct",
    "distribution_share_pct",
    "maps_to",
    "pref_enabled",
    "pref_rate",
    "pref_type",
    "pref_compounding",
    "pref_distribution_frequency",
    "pref_current_portion",
    "pref_accrued_portion",
    "deal_target_raise",
    "bucket_name",
    "load_notes",
]

HURDLE_HEADERS = [
    "deal_name",
    "deal_id",
    "class_name",
    "tier_order",
    "hurdle_name",
    "hurdle_rate_pct",
    "lp_pct",
    "gp_pct",
    "preferred_return_type",
    "is_final",
    "load_notes",
]

FINAL_HEADERS = [
    "deal_name",
    "deal_id",
    "class_name",
    "final_lp_pct",
    "final_gp_pct",
    "hurdle_name",
    "load_notes",
]


@dataclass
class ClassRow:
    name: str
    subscription_type: str
    source_class_name: str = ""
    status: str = "active"
    display_order: int = 0
    entity_name: str = ""
    start_date: str = ""
    offering_size: str = ""
    raise_amount_distributions: str = ""
    billing_raise_quota: str = ""
    minimum_investment: str = ""
    number_of_units: str = ""
    price_per_unit: str = ""
    visibility: str = ""
    equity_pct: str = ""
    distribution_share_pct: str = ""
    maps_to: str = ""
    pref_enabled: str = "no"
    pref_rate: str = ""
    pref_type: str = "single"
    pref_compounding: str = "simple"
    pref_distribution_frequency: str = "quarterly"
    pref_current_portion: str = ""
    pref_accrued_portion: str = ""
    bucket_name: str = ""
    load_notes: str = ""

    def to_sheet_row(self, deal_name: str, deal_id: str, deal_target_raise: str) -> list:
        return [
            deal_name,
            deal_id,
            self.source_class_name or self.name,
            self.name,
            self.subscription_type,
            self.status,
            self.display_order,
            self.entity_name,
            self.start_date,
            self.offering_size,
            self.raise_amount_distributions,
            self.billing_raise_quota,
            self.minimum_investment,
            self.number_of_units,
            self.price_per_unit,
            self.visibility,
            self.equity_pct,
            self.distribution_share_pct,
            self.maps_to or self.name,
            self.pref_enabled,
            self.pref_rate,
            self.pref_type,
            self.pref_compounding,
            self.pref_distribution_frequency,
            self.pref_current_portion,
            self.pref_accrued_portion,
            deal_target_raise,
            self.bucket_name,
            self.load_notes,
        ]


@dataclass
class HurdleRow:
    class_name: str
    tier_order: int = 1
    hurdle_name: str = "Hurdle 1 - Pref"
    hurdle_rate_pct: str = ""
    lp_pct: str = "100"
    gp_pct: str = "0"
    preferred_return_type: str = "cash_on_cash"
    is_final: str = "no"
    load_notes: str = ""

    def to_sheet_row(self, deal_name: str, deal_id: str) -> list:
        return [
            deal_name,
            deal_id,
            self.class_name,
            self.tier_order,
            self.hurdle_name,
            self.hurdle_rate_pct,
            self.lp_pct,
            self.gp_pct,
            self.preferred_return_type,
            self.is_final,
            self.load_notes,
        ]


@dataclass
class FinalTierRow:
    class_name: str
    final_lp_pct: str
    final_gp_pct: str
    hurdle_name: str = "Final"
    load_notes: str = ""

    def to_sheet_row(self, deal_name: str, deal_id: str) -> list:
        return [
            deal_name,
            deal_id,
            self.class_name,
            self.final_lp_pct,
            self.final_gp_pct,
            self.hurdle_name,
            self.load_notes,
        ]


@dataclass
class DealClasses:
    deal_name: str
    deal_target_raise: str = ""
    deal_id: str = ""
    classes: list[ClassRow] = field(default_factory=list)
    hurdles: list[HurdleRow] = field(default_factory=list)
    final_tiers: list[FinalTierRow] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "deal_name": self.deal_name,
            "deal_target_raise": self.deal_target_raise,
            "deal_id": self.deal_id,
            "classes": [c.__dict__ for c in self.classes],
            "hurdles": [h.__dict__ for h in self.hurdles],
            "final_tiers": [f.__dict__ for f in self.final_tiers],
        }

    @classmethod
    def from_dict(cls, data: dict) -> DealClasses:
        return cls(
            deal_name=str(data.get("deal_name") or ""),
            deal_target_raise=str(data.get("deal_target_raise") or ""),
            deal_id=str(data.get("deal_id") or ""),
            classes=[ClassRow(**c) for c in data.get("classes") or []],
            hurdles=[HurdleRow(**h) for h in data.get("hurdles") or []],
            final_tiers=[FinalTierRow(**f) for f in data.get("final_tiers") or []],
        )


def style_header(ws, cols: int) -> None:
    ws.row_dimensions[1].height = 40
    for c in range(1, cols + 1):
        cell = ws.cell(1, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def mark(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = example_fill
        cell.border = thin


def autosize(ws, max_w: int = 30) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        w = 12
        for cell in col:
            if cell.value is not None:
                w = max(w, min(max_w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = w


def write_workbook(data: DealClasses, out: Path, *, with_data: bool = True) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Classes"
    ws.append(CLASS_HEADERS)
    style_header(ws, len(CLASS_HEADERS))
    if with_data:
        for i, row in enumerate(data.classes, start=2):
            ws.append(row.to_sheet_row(data.deal_name, data.deal_id, data.deal_target_raise))
            mark(ws, i, len(CLASS_HEADERS))
    autosize(ws)
    ws.freeze_panes = "D2"

    ws = wb.create_sheet("Hurdles")
    ws.append(HURDLE_HEADERS)
    style_header(ws, len(HURDLE_HEADERS))
    if with_data:
        for i, row in enumerate(data.hurdles, start=2):
            ws.append(row.to_sheet_row(data.deal_name, data.deal_id))
            mark(ws, i, len(HURDLE_HEADERS))
    autosize(ws)
    ws.freeze_panes = "D2"

    ws = wb.create_sheet("Final Tier")
    ws.append(FINAL_HEADERS)
    style_header(ws, len(FINAL_HEADERS))
    if with_data:
        for i, row in enumerate(data.final_tiers, start=2):
            ws.append(row.to_sheet_row(data.deal_name, data.deal_id))
            mark(ws, i, len(FINAL_HEADERS))
    autosize(ws)
    ws.freeze_panes = "D2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.close()


def write_blank_template(out: Path) -> None:
    write_workbook(DealClasses(deal_name=""), out, with_data=False)
